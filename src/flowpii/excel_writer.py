from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import load_workbook

from .models import BifGraph, InventoryRow, Metadata

# Default template: Korea demo has reference sheets + validations
DEFAULT_TEMPLATE = (
    Path(__file__).resolve().parents[2]
    / "output"
    / "個人資料盤點清冊 韓國DEMO.xlsx"
)

COL = {
    "A": 1,
    "B": 2,
    "G": 7,
    "H": 8,
    "AF": 32,
    "AG": 33,
    "AH": 34,
}


def format_header(meta: Metadata) -> tuple[str, str]:
    date = meta.inventory_date or ""
    inv_ver = meta.inventory_version or ""
    bif_ver = meta.bif_version or ""
    # Prefer sample style when versions present
    left_parts = []
    if date or inv_ver or bif_ver:
        ver_bit = f" {inv_ver}".rstrip() if inv_ver else ""
        bif_bit = f" (BIF V{bif_ver})" if bif_ver else ""
        left_parts.append(f"盤點日期/版次：{date}{ver_bit}{bif_bit}".strip())
    left = left_parts[0] if left_parts else "盤點日期/版次："
    right = f"盤點部門：{meta.department}" if meta.department else "盤點部門："
    return left, right


def clear_data_rows(ws, start_row: int = 5) -> None:
    max_row = ws.max_row or start_row
    if max_row < start_row:
        return
    for r in range(start_row, max_row + 1):
        for c in range(1, (ws.max_column or 43) + 1):
            ws.cell(r, c).value = None


def write_inventory_excel(
    rows: list[InventoryRow],
    output_path: str | Path,
    *,
    metadata: Metadata | None = None,
    template_path: str | Path | None = None,
    graph: BifGraph | None = None,
) -> Path:
    """Copy company template and fill auto-derived columns only."""
    template = Path(template_path) if template_path else DEFAULT_TEMPLATE
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, out)

    wb = load_workbook(out)
    ws = wb["個資盤點清冊"] if "個資盤點清冊" in wb.sheetnames else wb.active

    meta = metadata or (graph.metadata if graph else Metadata())
    left, right = format_header(meta)
    ws["A1"] = left
    ws["G1"] = right

    clear_data_rows(ws, 5)

    for i, row in enumerate(rows):
        r = 5 + i
        d = row.as_dict()
        ws.cell(r, COL["A"], d["A"])
        ws.cell(r, COL["B"], d["B"])
        ws.cell(r, COL["G"], d["G"])
        ws.cell(r, COL["H"], d["H"])
        ws.cell(r, COL["AF"], d["AF"])
        ws.cell(r, COL["AG"], d["AG"])
        ws.cell(r, COL["AH"], d["AH"])

    wb.save(out)
    return out
