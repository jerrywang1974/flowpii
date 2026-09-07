from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from flowpii.excel_writer import write_inventory_excel
from flowpii.expand import expand_graph
from flowpii.models import BifGraph
from flowpii.normalize import row_match_key

ROOT = Path(__file__).resolve().parents[1]


def test_write_demo_excel(tmp_path: Path):
    graph = BifGraph.model_validate(
        json.loads((ROOT / "tests/fixtures/demo_graph.json").read_text(encoding="utf-8"))
    )
    rows = expand_graph(graph)
    out = tmp_path / "demo.xlsx"
    write_inventory_excel(rows, out, graph=graph)
    wb = load_workbook(out, data_only=True)
    ws = wb["個資盤點清冊"]
    got = []
    for r in range(5, ws.max_row + 1):
        g = ws.cell(r, 7).value
        if not g:
            continue
        got.append(
            {
                "A": ws.cell(r, 1).value,
                "B": ws.cell(r, 2).value,
                "G": g,
                "H": ws.cell(r, 8).value,
                "AF": ws.cell(r, 32).value,
                "AG": ws.cell(r, 33).value,
                "AH": ws.cell(r, 34).value,
            }
        )
    expected = json.loads((ROOT / "tests/golden/demo.json").read_text(encoding="utf-8"))[
        "rows"
    ]
    assert sorted(row_match_key(r) for r in got) == sorted(
        row_match_key(r) for r in expected
    )
